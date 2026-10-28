"""Build the platform's park dataset from the Earthwise Institute captive coal database.

Reads the unit-level sheet of the Earthwise workbook, rolls units up into sites (one per
named industrial park / complex, plus one per standalone captive plant), and writes:

  assets/js/parks.js        window.PARKS for the static site (map, filters, popups)
  assets/js/site-details.js window.SITE_DETAILS, per-site detail for the facility page
  data/parks.csv            the same sites as a flat table
  data/model_crosswalk.csv  each capacity-expansion model node -> platform site

The crosswalk is how model outputs get joined to sites. Model nodes are matched to the
nearest site automatically; data/crosswalk_overrides.csv pins the matches that need a human
call (wrong node coordinates, parks the model split or duplicated, etc.).

Usage:
  python3 scripts/build_parks.py EARTHWISE_XLSX [--extended EXT_XLSX] [--model-data DIR]

  EARTHWISE_XLSX   Earthwise-Institute_*_Indonesia-Captive-Coal-Database_*.xlsx
  --extended       Earthwise-Institute_*_Captive-Coal-Extended-Dims_*.xlsx (optional)
  --model-data     captive-indonesia-2025/data_indonesia/2030 (default: sibling checkout)
"""
import argparse
import csv
import glob
import json
import math
import os
import re
from collections import Counter, OrderedDict, defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = 'Indonesia captive coal plants'

# Earthwise status -> platform bucket. Cancelled, shelved, retired, closed and NA units are
# dropped: they are not capacity anyone has to transition.
STATUS = {'Operating': 'op', 'Construction': 'con', 'Pre-Construction': 'plan', 'Announced': 'plan'}

# Earthwise commodity -> one of the platform's eight sectors.
SECTOR = {
    'Nickel': 'nickel',
    'Aluminium': 'aluminium',
    'Pulp & Paper': 'pulp', 'Wood': 'pulp',
    'Cement & Building': 'cement',
    'Chemical & Petrochemical': 'chem', 'Chemical & Petrochemical, Iron & Steel': 'chem',
    'Foods': 'agro', 'Cpo': 'agro',
    'Copper, Gold & Silver': 'metals', 'Coal': 'metals', 'Iron & Steel': 'metals', 'Lithium': 'metals',
    'Textiles': 'mixed', 'PV & PV Parts': 'mixed', 'Industrial Park': 'mixed',
}

REGION = {
    'Sulawesi': ['Central Sulawesi', 'Southeast Sulawesi', 'South Sulawesi', 'North Sulawesi',
                 'Gorontalo', 'West Sulawesi'],
    'Maluku': ['North Maluku', 'Maluku'],
    'Kalimantan': ['West Kalimantan', 'Central Kalimantan', 'South Kalimantan', 'East Kalimantan',
                   'North Kalimantan'],
    'Sumatra': ['Aceh', 'North Sumatra', 'West Sumatra', 'Riau', 'Riau Islands', 'Jambi',
                'South Sumatra', 'Bengkulu', 'Lampung', 'Bangka Belitung Islands'],
    'Java': ['Banten', 'Jakarta', 'West Java', 'Central Java', 'Yogyakarta', 'East Java', 'Bali'],
    'Nusa Tenggara': ['West Nusa Tenggara', 'East Nusa Tenggara'],
    'Papua': ['Papua', 'West Papua', 'Central Papua', 'South Papua', 'Highland Papua',
              'Southwest Papua'],
}
PROV_REGION = {p: r for r, ps in REGION.items() for p in ps}

# Matches farther than this need an override (or are left unmatched).
AUTO_MATCH_KM = 15


def num(x):
    if x in (None, '', 'NA'):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    return float(re.sub(r'[^0-9.]', '', str(x)) or 0)


def year(x):
    try:
        return int(x)
    except (TypeError, ValueError):
        return None


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')


def park_short(name):
    """'IMIP Indonesia Morowali Industrial Park' -> ('IMIP', 'Indonesia Morowali Industrial Park').
    The abbreviation may also come second: 'Delong VDNIP Virtue Dragon ...' -> 'VDNIP'."""
    words = name.split(' ')
    for i in (0, 1):
        if len(words) > i + 1 and re.fullmatch(r'[A-Z]{2,6}', words[i]):
            return words[i], ' '.join(words[:i] + words[i + 1:])
    return None, name


def plant_short(name):
    return re.sub(r'\s+(power station|power plant)$', '', name.strip(), flags=re.I)


def km(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 6371 * 2 * math.asin(math.sqrt(h))


def read_sheet(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    head = next(rows)
    clean = lambda v: v.strip() if isinstance(v, str) else v
    return [{h: clean(v) for h, v in zip(head, r) if h} for r in rows if any(r)]


def read_units(wb):
    units = read_sheet(wb, SHEET)

    # Source errors we have confirmed, fixed field by field (data/earthwise_corrections.csv).
    by_id = {u['EW Project ID']: u for u in units}
    for c in csv.DictReader(open(os.path.join(ROOT, 'data', 'earthwise_corrections.csv'), encoding='utf-8')):
        by_id[c['ew_id']][c['field']] = c['value']
    return units


def site_key(u):
    park = u['Industrial Park / Complex'] or 'F'
    return ('park', park) if park != 'F' else ('standalone', u['Project Name English'])


def build_sites(units):
    groups = OrderedDict()
    for u in units:
        bucket = STATUS.get(u['Status'])
        if not bucket:
            continue
        groups.setdefault(site_key(u), []).append(dict(u, _bucket=bucket, _mw=num(u['Capacity (MW)'])))

    sites, ids = [], set()
    for (kind, name), us in groups.items():
        mw = Counter()
        for u in us:
            mw[u['_bucket']] += u['_mw']
        sector_mw = Counter()
        for u in us:
            sector_mw[SECTOR[u['Commodity'].strip('"')]] += u['_mw']
        anchor = max(us, key=lambda u: u['_mw'])  # site location = its largest live unit

        if kind == 'park':
            short, long_name = park_short(name)
            display = name if not short else '%s — %s' % (short, long_name)
            sid = slug(short or name)
        else:
            short = plant_short(name)
            display = short
            sid = slug(short)
        base, i = sid, 2
        while sid in ids:
            sid, i = '%s-%d' % (base, i), i + 1
        ids.add(sid)

        plants = OrderedDict()
        for u in us:
            p = plants.setdefault(u['Project Name English'].strip(), {
                'n': plant_short(u['Project Name English']), 'op': 0, 'con': 0, 'plan': 0, 'units': 0,
                'owner': u['Majority or Whole Ultimate Owner (Abbreviation)'] or u['Majority or Whole Ultimate Owner'],
                'ctry': u['Majority or Whole Ultimate Owner Country'], 'start': None,
                'com': u['Commodity'].strip('"')})
            p[u['_bucket']] += u['_mw']
            p['units'] += 1
            y = year(u['(Expected) Start year'])
            if y and (p['start'] is None or y < p['start']):
                p['start'] = y

        owner_mw, ctry_mw = Counter(), Counter()
        for u in us:
            owner_mw[u['Majority or Whole Ultimate Owner (Abbreviation)'] or u['Majority or Whole Ultimate Owner']] += u['_mw']
            ctry_mw[u['Majority or Whole Ultimate Owner Country']] += u['_mw']
        fin = set()
        for u in us:
            fin.update(t.strip() for t in str(u['Financing Transaction'] or '').split(';') if t.strip().startswith('EWTRANS'))
        starts = [year(u['(Expected) Start year']) for u in us if year(u['(Expected) Start year'])]
        prov = anchor['Province (Admin 1, English)']

        sites.append({
            'id': sid, 'name': display, 'short': short or display, 'kind': kind,
            'sector': sector_mw.most_common(1)[0][0],
            'commodities': sorted({u['Commodity'].strip('"') for u in us}),
            'prov': prov, 'region': PROV_REGION[prov], 'regency': anchor['Regency / City (Admin 2)'],
            'lat': round(float(anchor['Latitude']), 5), 'lng': round(float(anchor['Longitude']), 5),
            'loc': 'approximate' if any(u['Location Accuracy'] != 'Exact' for u in us) else 'exact',
            'op': round(mw['op']), 'con': round(mw['con']), 'plan': round(mw['plan']),
            'units': len(us),
            'owner': owner_mw.most_common(1)[0][0],
            'ownerCountry': ctry_mw.most_common(1)[0][0],
            'firstYear': min(starts) if starts else None,
            'financing': sorted(fin),
            'plants': [dict(p, op=round(p['op']), con=round(p['con']), plan=round(p['plan']))
                       for p in plants.values()],
            'ew': sorted(u['EW Project ID'] for u in us),
            '_key': (kind, name),
        })
    sites.sort(key=lambda s: -(s['op'] + s['con'] + s['plan']))
    return sites


def split_list(v, sep=';'):
    return [x.strip() for x in str(v or '').split(sep) if x.strip() and x.strip() not in ('NA', 'None', 'F')]


def build_details(units, sites, financing, extended):
    """Per-site detail for the facility page: every unit (including cancelled / shelved / retired
    ones at the same site), ownership, financing transactions and the industrial projects the
    site's power supports."""
    by_key = {s['_key']: s for s in sites}
    det, unit_site, plant_site = {}, {}, {}
    for s in sites:
        det[s['id']] = {'units': [], 'ultimate': [], 'countries': [], 'ownership': [], 'jv': False,
                        'epc': [], 'chinaRole': [], 'aliases': [], 'district': None, 'village': None,
                        'island': None, 'smallIsland': False, 'sources': [], 'finance': [], 'projects': [],
                        'otherPower': None}

    def add(lst, *vals):
        for v in vals:
            if v and v not in lst:
                lst.append(v)

    for u in units:
        s = by_key.get(site_key(u))
        if not s:
            continue
        d = det[s['id']]
        unit_site[u['EW Project ID']] = s['id']
        plant_site[u['Project Name English']] = s['id']
        d['units'].append({
            'plant': plant_short(u['Project Name English']), 'phase': u['Phase / Unit'],
            'mw': num(u['Capacity (MW)']), 'status': u['Status'],
            'start': year(u['(Expected) Start year']), 'retire': year(u['(Expected) Retire Year']),
            'owner': u['Owner'], 'epc': u['EPC'], 'project': u['Supported Project'],
            'fuel': u['Alternate Fuel'] if u['Alternate Fuel'] not in ('NA', 'None') else None,
        })
        add(d['ultimate'], u['Ultimate Owner (All)'])
        add(d['countries'], *split_list(u['Ultimate Owner Country (All)'], ','))
        add(d['ownership'], u['Ownership Type'])
        d['jv'] = d['jv'] or u['Solely Owned / JV'] == 'JV'
        add(d['epc'], *split_list(u['EPC']))
        add(d['chinaRole'], *split_list(u['Role of Chinese Entities']))
        add(d['aliases'], *split_list(u['Power Plant Name English Other']), u['Project Name Other Language'])
        if STATUS.get(u['Status']) and not d['district']:
            d['district'] = u['District / Subdistrict (Admin 3, English)']
            d['village'] = u['Village / Urban Village or Equivalent (Admin4, English)']
            d['island'] = u['Island']
            d['smallIsland'] = u['Major / Small island'] == 'Small Island'
        link = u['Primary Source Link (GEM Wiki)']
        if link and str(link).startswith('http'):
            add(d['sources'], link)

    for d in det.values():
        d['units'].sort(key=lambda x: (x['plant'], x['start'] or 9999, str(x['phase'])))

    for t in financing:
        ids = {unit_site[p] for p in split_list(t['Project ID']) if p in unit_site}
        refs = [t[k] for k in ('Ref1', 'Ref2', 'Ref3', 'Ref4') if t.get(k) and str(t[k]).startswith('http')]
        for sid in ids:
            det[sid]['finance'].append({
                'id': t['Transaction ID'], 'year': year(t['Financing Close Year']),
                'usd': num(t['Financing Amount USD Million Equivalent']) or None,
                'financiers': t['Financier / Issuer'], 'hq': sorted(set(split_list(t['Bank HQ']))),
                'type': ' · '.join(x for x in (t['Project / Corporate Finance'], t['Financing Strcture']) if x),
                'term': year(t['Financing Term (Years)']),
                'recipient': t['Financing Recepient'], 'phase': t['Financed Phase'],
                'captive': t['Financing specified for captive coal'] == 'T', 'refs': refs,
            })
    for d in det.values():
        d['finance'].sort(key=lambda f: f['year'] or 0)

    if extended:
        park_by_abbr = {s['short']: s['id'] for s in sites if s['kind'] == 'park'}
        park_names = [(s['_key'][1], s['id']) for s in sites if s['kind'] == 'park']

        def match(loc, plant):
            loc = str(loc or '')
            m = re.match(r'[A-Z]{2,6}\b', loc)
            if m and m.group(0) in park_by_abbr:
                return park_by_abbr[m.group(0)]
            for name, sid in park_names:
                if loc.startswith(name):
                    return sid
            return plant_site.get(plant)

        specs = [  # sheet, capacity column, unit, plant column, other-power column
            ('power source_Nickel project(aft', 'Capacity (ktpa)', 'ktpa', 'Captive Coal Info Availability', 'Other Captive Power'),
            ('power source_Alu projects(all)', 'Capacity (tmtpa)', 'ktpa', None, 'Other Captive Power'),
            ('power source_new downstream pro', 'Capacity', '', None, 'Other Power & Captive Power'),
        ]
        for sheet, capcol, unit, plantcol, othercol in specs:
            for p in read_sheet(extended, sheet):
                sid = match(p.get('Location'), p.get(plantcol) if plantcol else None)
                if not sid or not p.get('Project Name English'):
                    continue
                cap = p.get(capcol)
                det[sid]['projects'].append({
                    'name': p['Project Name English'], 'phase': p.get('Phase'),
                    'product': p.get('Product') or p.get('Product / Unit'),
                    'capacity': ('%s %s' % (cap, unit)).strip() if cap not in (None, 'NA') else None,
                    'status': p.get('Status'), 'start': p.get('(Expected) Start Year'),
                    'owner': p.get('Ultimat Owner'),
                    'otherPower': p.get(othercol) if p.get(othercol) not in (None, 'NA') else None,
                })
        for p in read_sheet(extended, 'power source_Industrial parks  '):
            sid = dict(park_names).get(p['Industrial Park / Complex Name'])
            if sid and p.get('Other Power Source Summary') not in (None, 'NA'):
                det[sid]['otherPower'] = p['Other Power Source Summary']
    return det


def read_model_nodes(model_dir):
    nodes = OrderedDict()
    for f in sorted(glob.glob(os.path.join(model_dir, '*', 'ip_generators.csv'))):
        island = os.path.basename(os.path.dirname(f))
        for r in csv.DictReader(open(f, encoding='utf-8-sig')):
            if r['New_Build'] == '1':
                continue
            n = nodes.setdefault((island, int(r['Industrial_Park'])), {
                'province': r['Province'], 'park': set(), 'plants': set(), 'mw': 0.0, 'coords': None})
            n['mw'] += num(r['Existing_Cap_MW'])
            if r['industrial_park'] not in ('', 'no'):
                n['park'].add(r['industrial_park'])
            if r['Resource_Other_Name']:
                n['plants'].add(r['Resource_Other_Name'])
            if r['latitude'] and n['coords'] is None:
                n['coords'] = (float(r['latitude']), float(r['longitude']))
    return nodes


def read_overrides(path):
    if not os.path.exists(path):
        return {}
    out = {}
    for r in csv.DictReader(open(path, encoding='utf-8')):
        out[(r['island'], int(r['ip']))] = (r['site_id'], r['note'])
    return out


def crosswalk(nodes, sites, overrides):
    by_id = {s['id']: s for s in sites}
    rows = []
    for (island, ip), n in nodes.items():
        best = min(sites, key=lambda s: km(n['coords'], (s['lat'], s['lng'])))
        dist = km(n['coords'], (best['lat'], best['lng']))
        if (island, ip) in overrides:
            sid, note = overrides[(island, ip)]
            method = 'override'
            if sid and sid not in by_id:
                raise SystemExit('crosswalk override %s:%d points at unknown site %r' % (island, ip, sid))
        elif dist <= AUTO_MATCH_KM:
            sid, note, method = best['id'], '', 'nearest'
        else:
            sid, note, method = '', 'no site within %d km' % AUTO_MATCH_KM, 'unmatched'
        rows.append({
            'island': island, 'ip': ip, 'model_province': n['province'],
            'model_park': ';'.join(sorted(n['park'])), 'model_plants': ';'.join(sorted(n['plants'])),
            'model_mw': round(n['mw']), 'site_id': sid, 'method': method,
            'nearest_site': best['id'], 'nearest_km': round(dist, 1), 'note': note,
        })
        if sid:
            by_id[sid].setdefault('model', []).append('%s:%d' % (island, ip))
    for s in sites:
        s.setdefault('model', [])
    return rows


def write_csv(path, rows, fields):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator='\n')
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('earthwise')
    ap.add_argument('--extended', help='Earthwise "Captive Coal, Extended Dimensions" workbook (industrial projects)')
    ap.add_argument('--model-data', default=os.path.join(ROOT, '..', 'captive-indonesia-2025', 'data_indonesia', '2030'))
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.earthwise, read_only=True, data_only=True)
    units = read_units(wb)
    sites = build_sites(units)
    ext = openpyxl.load_workbook(a.extended, read_only=True, data_only=True) if a.extended else None
    details = build_details(units, sites, read_sheet(wb, 'Captive coal debt financing tra'), ext)
    for s in sites:
        del s['_key']
    nodes = read_model_nodes(a.model_data)
    xw = crosswalk(nodes, sites, read_overrides(os.path.join(ROOT, 'data', 'crosswalk_overrides.csv')))

    js = ('/* GENERATED by scripts/build_parks.py from the Earthwise Institute Indonesian captive coal\n'
          '   database (%s). Do not edit by hand. Capacity is MW of captive coal by status:\n'
          '   op = operating, con = under construction, plan = announced / pre-construction.\n'
          '   model = capacity-expansion model nodes ("island:index") mapped to this site. */\n'
          'window.PARKS = [\n%s\n];\n') % (os.path.basename(a.earthwise),
                                          ',\n'.join(json.dumps(s, ensure_ascii=False) for s in sites))
    with open(os.path.join(ROOT, 'assets', 'js', 'parks.js'), 'w', encoding='utf-8') as f:
        f.write(js)

    js = ('/* GENERATED by scripts/build_parks.py: per-site detail for the facility page (units incl.\n'
          '   cancelled / shelved / retired, ownership, financing, supported industrial projects). */\n'
          'window.SITE_DETAILS = {\n%s\n};\n') % ',\n'.join(
              '%s: %s' % (json.dumps(k), json.dumps(v, ensure_ascii=False)) for k, v in details.items())
    with open(os.path.join(ROOT, 'assets', 'js', 'site-details.js'), 'w', encoding='utf-8') as f:
        f.write(js)

    flat = ['id', 'name', 'kind', 'sector', 'prov', 'region', 'regency', 'lat', 'lng', 'loc',
            'op', 'con', 'plan', 'units', 'owner', 'ownerCountry', 'firstYear', 'model']
    write_csv(os.path.join(ROOT, 'data', 'parks.csv'),
              [dict({k: s[k] for k in flat}, model=';'.join(s['model'])) for s in sites], flat)
    write_csv(os.path.join(ROOT, 'data', 'model_crosswalk.csv'), xw, list(xw[0]))

    tot = Counter()
    for s in sites:
        tot.update({k: s[k] for k in ('op', 'con', 'plan')})
    print('%d sites (%d parks, %d standalone) · operating %d MW · construction %d MW · planned %d MW'
          % (len(sites), sum(s['kind'] == 'park' for s in sites), sum(s['kind'] == 'standalone' for s in sites),
             tot['op'], tot['con'], tot['plan']))
    print('model nodes: %d · %s · sites with model results: %d'
          % (len(xw), dict(Counter(r['method'] for r in xw)), sum(bool(s['model']) for s in sites)))


if __name__ == '__main__':
    main()
